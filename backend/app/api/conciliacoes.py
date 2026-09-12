from __future__ import annotations

import io
from collections import Counter
from datetime import datetime, timezone
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import PlainTextResponse, StreamingResponse
from sqlalchemy.orm import Session

from app.api.auth import require_api_key
from app.database.deps import get_session
from app.models.audit_log import AuditAction, AuditLog
from app.models.bank_transaction import BankTransaction
from app.models.erp_transaction import ERPTransaction
from app.models.manual_adjustment import ManualAdjustment, ManualAdjustmentReason
from app.models.reconciliation import Reconciliation, ReconciliationLifecycleStatus
from app.models.reconciliation_match import ReconciliationMatchRow
from app.schemas.reconciliation import (
    DailyReconciliationRow,
    DashboardOut,
    DashboardStatusCount,
    ManualAdjustmentCreate,
    ReconciliationCreate,
    ReconciliationHistoryItem,
    ReconciliationMatchDetailOut,
    ReconciliationMatchOut,
    ReconciliationOut,
    RecoveredTitleOut,
)
from app.services.reconciliation_service import run_and_persist_reconciliation
from app.services.recovery_service import RecoveredTitle, find_recoverable_titles

router = APIRouter(prefix="/conciliacoes", tags=["conciliacoes"])

RECONCILED_STATUSES = {
    "CONCILIADO", "CONCILIADO D+1", "CONCILIADO D+2",
    "CORTE DE COMPETÊNCIA", "CORTE DE COMPETÊNCIA (fim do período importado)",
    "CONCILIADO MANUALMENTE", "CONCILIADO (desconto)", "CONCILIADO (título descontado)",
}


def _get_reconciliation_or_404(session: Session, reconciliation_id: UUID) -> Reconciliation:
    reconciliation = session.get(Reconciliation, reconciliation_id)
    if reconciliation is None:
        raise HTTPException(status_code=404, detail="Conciliação não encontrada.")
    return reconciliation


def _compute_reconciled_pct(session: Session, reconciliation_id: UUID) -> float:
    counts = Counter(
        status for (status,) in
        session.query(ReconciliationMatchRow.status).filter_by(reconciliation_id=reconciliation_id).all()
    )
    total = sum(counts.values())
    if not total:
        return 0.0
    reconciled = sum(c for s, c in counts.items() if s in RECONCILED_STATUSES)
    return round(100 * reconciled / total, 2)


@router.delete("/{reconciliation_id}", status_code=204)
def delete_reconciliation(
    reconciliation_id: UUID, performed_by: str,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    """Item 39/19 — exclusão é permitida, mas nunca silenciosa: gera log de
    auditoria com o estado anterior, e é bloqueada para competências
    FECHADA (reabra antes se realmente precisar apagar)."""
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    if reconciliation.status == ReconciliationLifecycleStatus.FECHADA:
        raise HTTPException(
            status_code=409,
            detail="Esta competência está fechada. Reabra antes de excluir.",
        )

    match_ids = [
        row.id for row in
        session.query(ReconciliationMatchRow.id).filter_by(reconciliation_id=reconciliation.id).all()
    ]
    if match_ids:
        session.query(ManualAdjustment).filter(
            ManualAdjustment.reconciliation_match_id.in_(match_ids)
        ).delete(synchronize_session=False)
        session.query(ReconciliationMatchRow).filter(
            ReconciliationMatchRow.id.in_(match_ids)
        ).delete(synchronize_session=False)

    session.add(AuditLog(
        action=AuditAction.EXCLUSAO, entity_type="Reconciliation", entity_id=reconciliation.id,
        performed_by=performed_by,
        before={
            "competencia_year": reconciliation.competencia_year,
            "competencia_month": reconciliation.competencia_month,
            "status": reconciliation.status.value,
        },
    ))
    session.delete(reconciliation)
    session.commit()


@router.get("", response_model=list[ReconciliationHistoryItem])
def list_reconciliations(
    bank_account_id: UUID,
    session: Session = Depends(get_session),
    _=Depends(require_api_key),
):
    """Item 22 — histórico de competências de uma conta, com o percentual
    já calculado, para a tela de histórico do cliente não obrigar o
    usuário a reimportar tudo de novo só pra ver onde parou."""
    reconciliations = (
        session.query(Reconciliation)
        .filter_by(bank_account_id=bank_account_id)
        .order_by(Reconciliation.competencia_year.desc(), Reconciliation.competencia_month.desc())
        .all()
    )
    return [
        ReconciliationHistoryItem(
            id=r.id, competencia_year=r.competencia_year, competencia_month=r.competencia_month,
            status=r.status.value, reconciled_pct=_compute_reconciled_pct(session, r.id),
            created_at=r.created_at, closed_at=r.closed_at,
        )
        for r in reconciliations
    ]


@router.post("", response_model=ReconciliationOut, status_code=201)
def create_reconciliation(payload: ReconciliationCreate, session: Session = Depends(get_session), _=Depends(require_api_key)):
    reconciliation = Reconciliation(**payload.model_dump())
    session.add(reconciliation)
    session.flush()
    session.commit()
    return reconciliation


@router.post("/{reconciliation_id}/executar", response_model=list[ReconciliationMatchOut])
def executar_conciliacao(reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key)):
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    if reconciliation.status == ReconciliationLifecycleStatus.FECHADA:
        raise HTTPException(
            status_code=409,
            detail="Esta competência está fechada. Reabra antes de rodar o motor novamente.",
        )

    matches = run_and_persist_reconciliation(session, reconciliation_id)
    if reconciliation.status == ReconciliationLifecycleStatus.RASCUNHO:
        reconciliation.status = ReconciliationLifecycleStatus.EM_ANALISE
    session.flush()

    session.add(AuditLog(
        action=AuditAction.CONCILIACAO, entity_type="Reconciliation", entity_id=reconciliation.id,
        performed_by="sistema", context={"matches_gerados": len(matches)},
    ))
    session.commit()
    return matches


@router.get("/{reconciliation_id}", response_model=ReconciliationOut)
def get_reconciliation(reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key)):
    return _get_reconciliation_or_404(session, reconciliation_id)


@router.get("/{reconciliation_id}/dashboard", response_model=DashboardOut)
def dashboard(reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key)):
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    matches = session.query(ReconciliationMatchRow).filter_by(reconciliation_id=reconciliation.id).all()

    counts = Counter(m.status for m in matches)
    by_status = []
    for status, count in counts.items():
        amount = 0.0
        for m in matches:
            if m.status != status:
                continue
            if m.bank_transaction and m.bank_transaction.principal_amount is not None:
                amount += float(m.bank_transaction.client_amount or 0)
            elif m.erp_transaction and m.erp_transaction.incoming_amount is not None:
                amount += float(m.erp_transaction.incoming_amount)
        by_status.append(DashboardStatusCount(status=status, count=count, total_amount=round(amount, 2)))

    reconciled_count = sum(c.count for c in by_status if c.status in RECONCILED_STATUSES)
    reconciled_amount = sum(c.total_amount for c in by_status if c.status in RECONCILED_STATUSES)
    divergent_amount = sum(c.total_amount for c in by_status if c.status == "VALOR DIVERGENTE")

    # % de conciliação sobre o total de registros do motor. Título
    # descontado deixou de ser tratado como estruturalmente fora de escopo
    # (Regra 13 já explica a maioria via antecipação/CRP032A1) — o que
    # sobrar sem explicação entra no denominador como qualquer outro
    # residual (banco-sem-ERP, ERP-sem-banco), sem inflar nem esconder.
    total_matches = sum(c.count for c in by_status)

    bank_titles_count = session.query(BankTransaction).filter_by(
        bank_account_id=reconciliation.bank_account_id,
        import_file_id=reconciliation.bank_import_file_id,
    ).count()
    erp_titles_count = session.query(ERPTransaction).filter_by(
        bank_account_id=reconciliation.bank_account_id,
        import_file_id=reconciliation.erp_import_file_id,
    ).count()

    return DashboardOut(
        reconciliation_id=reconciliation.id,
        bank_titles_count=bank_titles_count,
        erp_titles_count=erp_titles_count,
        reconciled_count=reconciled_count,
        reconciled_pct=round(100 * reconciled_count / total_matches, 2) if total_matches else 0.0,
        reconciled_amount=round(reconciled_amount, 2),
        divergent_amount=round(divergent_amount, 2),
        by_status=by_status,
    )


@router.get("/{reconciliation_id}/titulos", response_model=list[ReconciliationMatchOut])
def listar_titulos(
    reconciliation_id: UUID,
    status: str | None = None,
    search: str | None = Query(default=None, description="Busca por título, cliente, valor, NF ou documento"),
    session: Session = Depends(get_session),
    _=Depends(require_api_key),
):
    query = session.query(ReconciliationMatchRow).filter_by(reconciliation_id=reconciliation_id)
    if status:
        query = query.filter_by(status=status)

    rows = query.all()

    if search:
        needle = search.strip().lower()

        def matches_search(m: ReconciliationMatchRow) -> bool:
            haystacks = []
            if m.bank_transaction:
                haystacks += [m.bank_transaction.seu_numero, m.bank_transaction.payer_name,
                              str(m.bank_transaction.principal_amount)]
            if m.erp_transaction:
                haystacks += [m.erp_transaction.invoice_number_raw, m.erp_transaction.nota_fiscal,
                              str(m.erp_transaction.incoming_amount)]
            return any(needle in str(h).lower() for h in haystacks if h)

        rows = [m for m in rows if matches_search(m)]

    return rows


@router.get("/{reconciliation_id}/titulos/{match_id}", response_model=ReconciliationMatchDetailOut)
def detalhe_titulo(
    reconciliation_id: UUID, match_id: UUID,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    match = session.get(ReconciliationMatchRow, match_id)
    if match is None or match.reconciliation_id != reconciliation_id:
        raise HTTPException(status_code=404, detail="Registro de conciliação não encontrado.")
    return match


@router.post("/{reconciliation_id}/titulos/{match_id}/ajuste-manual", response_model=ReconciliationMatchOut)
def ajuste_manual(
    reconciliation_id: UUID, match_id: UUID, payload: ManualAdjustmentCreate,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    if reconciliation.status == ReconciliationLifecycleStatus.FECHADA:
        raise HTTPException(status_code=409, detail="Competência fechada — não é possível ajustar. Reabra antes.")

    match = session.get(ReconciliationMatchRow, match_id)
    if match is None or match.reconciliation_id != reconciliation_id:
        raise HTTPException(status_code=404, detail="Registro de conciliação não encontrado.")

    try:
        reason = ManualAdjustmentReason(payload.reason)
    except ValueError:
        raise HTTPException(
            status_code=422,
            detail=f"Motivo inválido. Use um de: {', '.join(r.value for r in ManualAdjustmentReason)}",
        )

    previous_status = match.status
    session.add(ManualAdjustment(
        reconciliation_match_id=match.id, reason=reason, observation=payload.observation,
        previous_status=previous_status, new_status=payload.new_status,
        performed_by=payload.performed_by,
    ))
    match.status = payload.new_status
    match.is_manual_override = True
    session.flush()

    session.add(AuditLog(
        action=AuditAction.AJUSTE_MANUAL, entity_type="ReconciliationMatchRow", entity_id=match.id,
        performed_by=payload.performed_by,
        before={"status": previous_status}, after={"status": payload.new_status},
        context={"reason": reason.value, "observation": payload.observation},
    ))
    session.commit()
    return match


@router.get("/{reconciliation_id}/diaria", response_model=list[DailyReconciliationRow])
def conciliacao_diaria(reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key)):
    """Item 16 — compara o total liquidado no banco com o total baixado no
    ERP, dia a dia, para responder 'em qual dia ERP e banco deixaram de
    fechar'. Os totais são independentes do pareamento título a título
    (um título D+1 conta no dia em que o dinheiro efetivamente moveu em
    cada lado) — é isso que deixa visível um dia com diferença mesmo
    quando, título a título, tudo está com status CONCILIADO D+1."""
    from app.reconciliation.engine import SETTLEMENT_OPERATION_TYPES, ERP_TIPO_DOCUMENTO_IN_SCOPE

    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)

    bank_rows = session.query(BankTransaction).filter_by(
        import_file_id=reconciliation.bank_import_file_id,
        bank_account_id=reconciliation.bank_account_id,
    ).all()
    erp_rows = session.query(ERPTransaction).filter_by(
        import_file_id=reconciliation.erp_import_file_id,
        bank_account_id=reconciliation.bank_account_id,
    ).all()

    bank_by_day: dict = {}
    for b in bank_rows:
        if b.operation_type not in SETTLEMENT_OPERATION_TYPES:
            continue
        bank_by_day.setdefault(b.movement_date, 0.0)
        bank_by_day[b.movement_date] += float(b.client_amount or 0)

    erp_by_day: dict = {}
    for e in erp_rows:
        if e.tipo != "Rec. Dup." or e.is_anticipation or e.tipo_documento not in ERP_TIPO_DOCUMENTO_IN_SCOPE:
            continue
        erp_by_day.setdefault(e.transaction_date, 0.0)
        erp_by_day[e.transaction_date] += float(e.incoming_amount or 0)

    matches = session.query(ReconciliationMatchRow).filter_by(reconciliation_id=reconciliation_id).all()
    match_ids_by_day: dict = {}
    for m in matches:
        day = m.bank_transaction.movement_date if m.bank_transaction else (
            m.erp_transaction.transaction_date if m.erp_transaction else None
        )
        if day is None:
            continue
        match_ids_by_day.setdefault(day, []).append(m.id)

    all_days = sorted(set(bank_by_day) | set(erp_by_day))
    rows = []
    for day in all_days:
        bank_total = round(bank_by_day.get(day, 0.0), 2)
        erp_total = round(erp_by_day.get(day, 0.0), 2)
        diff = round(bank_total - erp_total, 2)
        rows.append(DailyReconciliationRow(
            date=day, bank_total=bank_total, erp_total=erp_total, difference=diff,
            status="OK" if abs(diff) < 0.01 else "DIVERGENTE",
            match_ids=match_ids_by_day.get(day, []),
        ))
    return rows


@router.get("/{reconciliation_id}/recuperacao", response_model=list[RecoveredTitleOut])
def listar_titulos_recuperaveis(
    reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key),
):
    """Cruza os 'TÍTULO DESCONTADO SEM CORRESPONDÊNCIA' com o FTP050
    (Relação de NF Emitidas) para descobrir Local e Cliente de documentos
    apagados por engano — ver app/services/recovery_service.py."""
    _get_reconciliation_or_404(session, reconciliation_id)
    try:
        return find_recoverable_titles(session, reconciliation_id)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


def _recovered_to_row(r: RecoveredTitle) -> dict:
    return {
        "Local": r.local_id if r.resolved else None,
        "Duplicata": r.nota_fiscal,
        "Sequência": r.sequencia,
        "Status": "Cobrança Ativa",
        "Cliente": r.cliente_id if r.resolved else None,
        "Representante": r.representante_id,
        "Data Emissão": r.data_emissao.strftime("%d/%m/%Y") if r.data_emissao else None,
        "Vencimento": r.due_date.strftime("%d/%m/%Y") if r.due_date else None,
        "Valor de Emissão": r.principal_amount,
        "Banco": 341,
        "Agência": r.agency,
        "Tipo Documento": "BOLETO BANCARIO [5]",
        "Forma Pagamento NF-e 4.0": "Boleto Bancário-[15]",
        "Número do Título": r.nosso_numero,
        "_Seu Número (banco)": r.seu_numero,
        "_Pagador (banco)": r.payer_name,
        "_Razão Social": r.razao_social,
        "_Representante (nome)": r.representante_nome,
        "_Local (nome)": r.local_nome,
        "_Data liquidação banco": r.movement_date.strftime("%d/%m/%Y"),
        "_Resolvido": "SIM" if r.resolved else "NÃO — revisar manualmente",
    }


@router.get("/{reconciliation_id}/recuperacao/excel")
def exportar_recuperacao_excel(
    reconciliation_id: UUID, session: Session = Depends(get_session), _=Depends(require_api_key),
):
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter

    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    titulos = find_recoverable_titles(session, reconciliation_id)

    FORM_COLS = [
        "Local", "Duplicata", "Sequência", "Status", "Cliente", "Representante",
        "Data Emissão", "Vencimento", "Valor de Emissão", "Banco", "Agência",
        "Tipo Documento", "Forma Pagamento NF-e 4.0", "Número do Título",
    ]
    REF_COLS = [
        "_Seu Número (banco)", "_Pagador (banco)", "_Razão Social",
        "_Representante (nome)", "_Local (nome)", "_Data liquidação banco", "_Resolvido",
    ]
    ALL_COLS = FORM_COLS + REF_COLS

    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos a reimputar"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="0F766E")
    ref_header_fill = PatternFill("solid", fgColor="78716C")
    cell_font = Font(name="Arial", size=10)
    ref_cell_font = Font(name="Arial", size=10, italic=True, color="57534E")
    thin = Side(style="thin", color="D6D3D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(ALL_COLS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.lstrip("_"))
        cell.font = header_font
        cell.fill = header_fill if col_name in FORM_COLS else ref_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    valor_col_idx = FORM_COLS.index("Valor de Emissão") + 1
    ws.cell(row=1, column=valor_col_idx).comment = Comment(
        "Extraído da coluna 'Valor Inicial (R$)' da Francesinha — a 'Valor Final (R$)' "
        "vem R$ 0,00 para título descontado (carteira de antecipação bancária). "
        "Confirmar antes de importar em massa.",
        "Sistema de Conciliação",
    )

    for row_idx, t in enumerate(titulos, start=2):
        row = _recovered_to_row(t)
        for col_idx, col_name in enumerate(ALL_COLS, start=1):
            value = row.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font if col_name in FORM_COLS else ref_cell_font
            cell.border = border
            if col_name == "Valor de Emissão" and value is not None:
                cell.number_format = "#,##0.00"
            if col_name == "_Resolvido" and value != "SIM":
                cell.font = Font(name="Arial", size=10, bold=True, color="B91C1C")

    widths = {
        "Local": 8, "Duplicata": 11, "Sequência": 10, "Status": 15, "Cliente": 9,
        "Representante": 12, "Data Emissão": 13, "Vencimento": 12, "Valor de Emissão": 15,
        "Banco": 8, "Agência": 9, "Tipo Documento": 20, "Forma Pagamento NF-e 4.0": 22,
        "Número do Título": 15, "_Seu Número (banco)": 16, "_Pagador (banco)": 30,
        "_Razão Social": 45, "_Representante (nome)": 30, "_Local (nome)": 16,
        "_Data liquidação banco": 16, "_Resolvido": 22,
    }
    for col_idx, col_name in enumerate(ALL_COLS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(col_name, 14)
    ws.freeze_panes = "A2"
    if titulos:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_COLS))}{len(titulos) + 1}"

    legend = wb.create_sheet("Legenda")
    legend.cell(row=1, column=1, value="Legenda").font = Font(name="Arial", bold=True, size=12)
    notas = [
        "",
        "Colunas VERDES: campos exatos do formulário do ERP (CRP015A1).",
        "Colunas CINZAS: dados de referência/auditoria — não fazem parte do formulário.",
        "'Local'/'Cliente' vazios e '_Resolvido' = NÃO: a NF não foi encontrada com confiança no FTP050/FTP021A1 — revisar manualmente antes de importar.",
        "Valor de Emissão vem de 'Valor Inicial (R$)' da Francesinha (ver comentário na célula do cabeçalho).",
        "Data Emissão vem do FTP050 (ou do FTP021A1 quando o FTP050 não resolveu).",
        "Representante vem do FTP021A1 — pode ficar vazio mesmo quando Local/Cliente foram resolvidos, se esse arquivo não tiver sido importado.",
    ]
    for i, nota in enumerate(notas, start=2):
        legend.cell(row=i, column=1, value=nota).font = Font(name="Arial", size=10)
    legend.column_dimensions["A"].width = 100

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"reimputacao_{reconciliation.competencia_year}_{reconciliation.competencia_month:02d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{reconciliation_id}/recuperacao/{match_id}/script", response_class=PlainTextResponse)
def gerar_script_console_por_titulo(
    reconciliation_id: UUID, match_id: UUID,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    """Gera o script de console (F12) para UM único documento, usando o
    template fornecido (com seletores reais validados pelo cliente para
    os campos de busca customizados — Local, Status, Tipo Documento,
    Forma Pagamento). Um botão 'Criar Console' por linha, não um arquivo
    combinado — o usuário copia, cola no console com o formulário de
    NOVO documento aberto, confere e salva, um de cada vez."""
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    titulos = find_recoverable_titles(session, reconciliation_id)
    titulo = next((t for t in titulos if t.match_id == match_id), None)
    if titulo is None:
        raise HTTPException(status_code=404, detail="Título não encontrado na recuperação desta conciliação.")

    account_number = reconciliation.bank_account.account_number if reconciliation.bank_account else ""
    local_display = f"{titulo.local_nome} [{titulo.local_id}]" if titulo.local_nome and titulo.local_id is not None else ""

    import json
    dados = {
        "seuNumero": titulo.seu_numero,
        "local": local_display,
        "status": "Cobrança Ativa",
        "cliente": titulo.cliente_id or "",
        "representante": titulo.representante_id or "",
        "dataEmissao": titulo.data_emissao.strftime("%d/%m/%Y") if titulo.data_emissao else "",
        "vencimento": titulo.due_date.strftime("%d/%m/%Y") if titulo.due_date else "",
        "valorEmissao": f"{titulo.principal_amount:.2f}".replace(".", ",") if titulo.principal_amount is not None else "",
        "banco": "341",
        "agencia": titulo.agency or "",
        "contaCorrente": account_number,
        "tipoDocumento": "5",
        "formaPagamento": "15",
        "numeroTitulo": titulo.nosso_numero,
    }
    dados_js = json.dumps(dados, ensure_ascii=False, indent=8)

    if not titulo.resolved:
        aviso = (
            "// ATENÇÃO: este documento NÃO foi resolvido com confiança (Local/Cliente\n"
            "// não encontrados no FTP050/FTP021A1) — confira 'local' e 'cliente' abaixo\n"
            "// antes de rodar, provavelmente estão vazios.\n"
        )
    else:
        aviso = ""

    script = f"""// Script gerado pelo Sistema de Conciliação — reimputação de documento apagado.
// Título {titulo.seu_numero} (NF {titulo.nota_fiscal}/{titulo.sequencia}) — {titulo.razao_social or titulo.payer_name or "cliente não identificado"}
// Cole no console (F12) com o formulário de NOVO documento aberto no CRP015A1.
{aviso}(async function () {{

    // ============================================================
    // DADOS PARA TESTE
    // ============================================================

    const DADOS = {dados_js};


    // ============================================================
    // CONFIGURAÇÃO DE VELOCIDADE
    // ============================================================

    const TEMPO = {{
        campo: 120,
        dropdownAbrir: 180,
        dropdownSelecionar: 180,
        estabilizar: 250,

        // Apenas Sequência precisa de um pouco mais,
        // pois sabemos que dispara find('CODCLIENTE')
        cargaSequencia: 800,

        // Quanto tempo observar rapidamente se um modal apareceu
        janelaModal: 700
    }};


    // ============================================================
    // UTILITÁRIOS
    // ============================================================

    const esperar = ms =>
        new Promise(resolve => setTimeout(resolve, ms));


    function getAllDocs() {{

        const docs = [document];

        function buscar(doc) {{

            doc.querySelectorAll('iframe').forEach(iframe => {{

                try {{

                    const d =
                        iframe.contentDocument ||
                        iframe.contentWindow.document;

                    if (d && !docs.includes(d)) {{

                        docs.push(d);
                        buscar(d);

                    }}

                }} catch (e) {{}}

            }});

        }}

        buscar(document);

        return docs;
    }}


    function localizar(name) {{

        for (const doc of getAllDocs()) {{

            const campo =
                doc.querySelector(`[name="${{name}}"]`);

            if (campo) {{
                return campo;
            }}

        }}

        return null;
    }}


    // ============================================================
    // MODAL ATENÇÃO
    // ============================================================

    function localizarModalAtencao() {{

        for (const doc of getAllDocs()) {{

            const modal =
                doc.querySelector(
                    '#DHTMLSuite_modalBox_contentDiv'
                );

            if (!modal) {{
                continue;
            }}


            const style =
                doc.defaultView?.getComputedStyle(modal);


            const visivel =
                style &&
                style.display !== 'none' &&
                style.visibility !== 'hidden';


            if (visivel) {{

                return {{
                    modal,
                    doc
                }};

            }}

        }}

        return null;
    }}


    // ============================================================
    // FECHAR MODAL IMEDIATAMENTE
    // ============================================================

    async function fecharModalAgora() {{

        const resultado =
            localizarModalAtencao();


        if (!resultado) {{
            return false;
        }}


        const {{
            modal,
            doc
        }} = resultado;


        const texto =
            (modal.innerText || '')
                .replace(/\\s+/g, ' ')
                .trim();


        console.warn(
            `⚠️ ERP: ${{texto}}`
        );


        const botao =
            modal.querySelector(
                'a[onclick*="closeMessage"]'
            );


        if (botao) {{

            botao.click();

        }} else {{

            try {{

                if (
                    typeof doc.defaultView.closeMessage ===
                    'function'
                ) {{

                    doc.defaultView.closeMessage();

                }}

            }} catch (e) {{}}

        }}


        console.log(
            '✅ Aviso fechado automaticamente.'
        );


        await esperar(100);

        return true;
    }}


    // ============================================================
    // OBSERVAR RAPIDAMENTE SE MODAL APARECE
    // ============================================================

    async function observarModal(
        tempo = TEMPO.janelaModal
    ) {{

        const inicio =
            Date.now();


        while (
            Date.now() - inicio < tempo
        ) {{

            if (localizarModalAtencao()) {{

                await fecharModalAgora();

                // Continua observando por pouco tempo,
                // caso outro modal apareça.

                await esperar(80);

            }} else {{

                await esperar(50);

            }}

        }}

    }}


    // ============================================================
    // PREENCHER INPUT
    // ============================================================

    async function preencher(
        name,
        valor,
        readonly = false
    ) {{

        // Fecha algum modal que já esteja aberto.
        await fecharModalAgora();


        const campo =
            localizar(name);


        if (!campo) {{

            console.error(
                `❌ Campo não encontrado: ${{name}}`
            );

            return false;
        }}


        const tinhaReadonly =
            campo.hasAttribute('readonly');


        if (
            readonly &&
            tinhaReadonly
        ) {{

            campo.removeAttribute('readonly');

        }}


        campo.focus();

        campo.value =
            valor;


        campo.dispatchEvent(
            new Event(
                'input',
                {{ bubbles: true }}
            )
        );


        campo.dispatchEvent(
            new Event(
                'change',
                {{ bubbles: true }}
            )
        );


        campo.dispatchEvent(
            new Event(
                'blur',
                {{ bubbles: true }}
            )
        );


        if (
            readonly &&
            tinhaReadonly
        ) {{

            campo.setAttribute(
                'readonly',
                ''
            );

        }}


        console.log(
            `✅ ${{name}}: ${{valor}}`
        );


        await esperar(
            TEMPO.campo
        );


        // Caso o onchange tenha criado modal,
        // fecha imediatamente se ele já apareceu.
        await fecharModalAgora();


        return true;
    }}


    // ============================================================
    // VALOR MONETÁRIO
    // ============================================================

    async function preencherValorMonetario(
        name,
        valor
    ) {{

        await fecharModalAgora();


        const campo =
            localizar(name);


        if (!campo) {{

            console.error(
                `❌ Campo não encontrado: ${{name}}`
            );

            return false;
        }}


        const valorInterno =
            String(valor)
                .replace(/\\./g, '')
                .replace(',', '.');


        campo.focus();

        campo.value =
            valorInterno;


        campo.dispatchEvent(
            new Event(
                'input',
                {{ bubbles: true }}
            )
        );


        campo.dispatchEvent(
            new Event(
                'change',
                {{ bubbles: true }}
            )
        );


        campo.dispatchEvent(
            new Event(
                'blur',
                {{ bubbles: true }}
            )
        );


        await esperar(
            TEMPO.campo
        );


        console.log(
            `✅ ${{name}}: ${{valor}} | ERP: ${{campo.value}}`
        );


        await fecharModalAgora();

        return true;
    }}


    // ============================================================
    // NORMALIZAR TEXTO
    // ============================================================

    function normalizar(texto) {{

        return String(texto || '')
            .replace(/\\s+/g, ' ')
            .replace(':', '')
            .trim()
            .toLowerCase();

    }}


    // ============================================================
    // LOCALIZAR DROPDOWN
    // ============================================================

    function localizarDropdown(
        rotulo
    ) {{

        const alvo =
            normalizar(rotulo);


        for (const doc of getAllDocs()) {{

            const elementos = [
                ...doc.querySelectorAll(
                    'label, div, span, td, th'
                )
            ];


            for (const elemento of elementos) {{

                const texto =
                    normalizar(
                        elemento.textContent
                    );


                if (
                    texto !== alvo &&
                    !texto.startsWith(alvo)
                ) {{

                    continue;

                }}


                let atual =
                    elemento;


                for (
                    let nivel = 0;
                    nivel < 6 && atual;
                    nivel++
                ) {{

                    const input =
                        atual.querySelector?.(
                            'input.search[autocomplete="off"][tabindex="0"]'
                        );


                    if (input) {{

                        return input;

                    }}


                    atual =
                        atual.parentElement;

                }}

            }}

        }}


        return null;
    }}


    // ============================================================
    // ABRIR DROPDOWN
    // ============================================================

    async function abrirDropdown(
        input
    ) {{

        await fecharModalAgora();


        input.focus();


        input.dispatchEvent(
            new MouseEvent(
                'mousedown',
                {{ bubbles: true }}
            )
        );


        input.click();


        input.dispatchEvent(
            new MouseEvent(
                'mouseup',
                {{ bubbles: true }}
            )
        );


        await esperar(
            TEMPO.dropdownAbrir
        );

    }}


    // ============================================================
    // DROPDOWN POR TEXTO
    // ============================================================

    async function selecionarPorTexto(
        rotulo,
        textoDesejado
    ) {{

        const input =
            localizarDropdown(rotulo);


        if (!input) {{

            console.error(
                `❌ Dropdown não encontrado: ${{rotulo}}`
            );

            return false;
        }}


        await abrirDropdown(input);


        const desejado =
            normalizar(textoDesejado);


        for (const doc of getAllDocs()) {{

            const itens = [
                ...doc.querySelectorAll(
                    'div.item'
                )
            ];


            const item =
                itens.find(el =>

                    el.offsetParent !== null &&

                    normalizar(
                        el.textContent
                    ) === desejado

                );


            if (item) {{

                item.click();


                console.log(
                    `✅ ${{rotulo}}: ${{textoDesejado}}`
                );


                await esperar(
                    TEMPO.dropdownSelecionar
                );


                await fecharModalAgora();

                return true;
            }}

        }}


        console.error(
            `❌ Opção não encontrada: ${{textoDesejado}}`
        );


        return false;
    }}


    // ============================================================
    // DROPDOWN POR DATA-VALUE
    // ============================================================

    async function selecionarPorValor(
        rotulo,
        valor
    ) {{

        const input =
            localizarDropdown(rotulo);


        if (!input) {{

            console.error(
                `❌ Dropdown não encontrado: ${{rotulo}}`
            );

            return false;
        }}


        await abrirDropdown(input);


        for (const doc of getAllDocs()) {{

            const itens = [
                ...doc.querySelectorAll(
                    'div.item[data-value]'
                )
            ];


            const item =
                itens.find(el =>

                    el.offsetParent !== null &&

                    el.getAttribute(
                        'data-value'
                    ) === String(valor)

                );


            if (item) {{

                const descricao =
                    (item.textContent || '')
                        .trim();


                item.click();


                console.log(
                    `✅ ${{rotulo}}: ${{descricao}}`
                );


                await esperar(
                    TEMPO.dropdownSelecionar
                );


                await fecharModalAgora();

                return true;
            }}

        }}


        console.error(
            `❌ ${{rotulo}}: valor ${{valor}} não encontrado`
        );


        return false;
    }}


    // ============================================================
    // INÍCIO
    // ============================================================

    console.log(
        '🚀 INICIANDO PREENCHIMENTO RÁPIDO'
    );


    // ============================================================
    // CALCULAR DUPLICATA / SEQUÊNCIA
    // ============================================================

    const numero =
        DADOS.seuNumero
            .replace(/\\D/g, '');


    const duplicata =
        numero.slice(0, -2);


    const sequencia =
        numero.slice(-2);


    console.log(
        `ℹ️ ${{numero}} → Duplicata ${{duplicata}} / Seq ${{sequencia}}`
    );


    // ============================================================
    // 1. DUPLICATA
    // ============================================================

    await preencher(
        'NUMFATURA',
        duplicata
    );


    // ============================================================
    // 2. SEQUÊNCIA
    // ============================================================

    await preencher(
        'NUMSEQUENCIA',
        sequencia
    );


    // Sequência dispara find('CODCLIENTE'), que tenta localizar um
    // documento existente com essa Duplicata+Sequência para puxar
    // Cliente/Banco/Agência/Tipo Documento dele. Como o documento original
    // foi apagado (é por isso que estamos recriando), o ERP mostra
    // 'Atenção: ... não encontrado' — ESPERADO neste cenário. Vigiamos
    // por uma janela maior aqui (em vez de só esperar um tempo fixo)
    // porque esse aviso pode demorar a aparecer, e se aparecer depois do
    // ponto em que checamos, ele trava a tela sem o script perceber.

    await observarModal(2000);


    await fecharModalAgora();


    // ============================================================
    // 3. LOCAL
    // ============================================================

    await selecionarPorTexto(
        'Local',
        DADOS.local
    );


    // ============================================================
    // 4. STATUS
    // ============================================================

    await selecionarPorTexto(
        'Status',
        DADOS.status
    );


    // ============================================================
    // 5. CLIENTE
    // ============================================================

    await preencher(
        'CODCLIENTE',
        DADOS.cliente
    );


    // Observa brevemente o lookup sem pausa longa.
    await observarModal(600);


    // ============================================================
    // 6. REPRESENTANTE
    // ============================================================

    await preencher(
        'REPRESENTANTE',
        DADOS.representante
    );


    await observarModal(500);


    // ============================================================
    // 7. DATA EMISSÃO
    // ============================================================

    await preencher(
        'DATAEMISSAO',
        DADOS.dataEmissao
    );


    // ============================================================
    // 8. VENCIMENTO
    // ============================================================

    await preencher(
        'DATAVENCTO',
        DADOS.vencimento
    );


    await esperar(200);

    await fecharModalAgora();


    // ============================================================
    // 9. VALOR EMISSÃO
    // ============================================================

    await preencherValorMonetario(
        'VALOREMISSAO',
        DADOS.valorEmissao
    );


    // ============================================================
    // 10. CONTA CORRENTE
    // ============================================================

    await preencher(
        'CONTACORRENTE',
        DADOS.contaCorrente
    );


    // Conta dispara refreshpag('NUMTITULO').
    // Observa rapidamente eventual aviso.
    await observarModal(700);


    // ============================================================
    // 11. BANCO
    // ============================================================

    await preencher(
        'CODBANCO',
        DADOS.banco,
        true
    );


    await observarModal(400);


    // ============================================================
    // 12. AGÊNCIA
    // ============================================================

    await preencher(
        'CODAGENCIA',
        DADOS.agencia,
        true
    );


    await observarModal(400);


    // ============================================================
    // 13. TIPO DOCUMENTO
    // ============================================================

    await selecionarPorValor(
        'Tipo Documento',
        DADOS.tipoDocumento
    );


    // ============================================================
    // 14. FORMA PAGAMENTO
    // ============================================================

    await selecionarPorValor(
        'Forma Pagamento',
        DADOS.formaPagamento
    );


    // ============================================================
    // 15. NÚMERO DO TÍTULO
    // ============================================================

    await preencher(
        'NUMTITULO',
        DADOS.numeroTitulo
    );


    // ============================================================
    // GARANTIA FINAL
    // ============================================================

    await fecharModalAgora();


    // ============================================================
    // FINAL
    // ============================================================

    console.log(
        '=========================================='
    );

    console.log(
        '🎉 PREENCHIMENTO FINALIZADO'
    );

    console.log(
        `Duplicata: ${{duplicata}} | Sequência: ${{sequencia}}`
    );

    console.log(
        `Local: ${{DADOS.local}}`
    );

    console.log(
        `Status: ${{DADOS.status}}`
    );

    console.log(
        `Cliente: ${{DADOS.cliente}}`
    );

    console.log(
        `Representante: ${{DADOS.representante}}`
    );

    console.log(
        `Emissão: ${{DADOS.dataEmissao}}`
    );

    console.log(
        `Vencimento: ${{DADOS.vencimento}}`
    );

    console.log(
        `Valor: ${{DADOS.valorEmissao}}`
    );

    console.log(
        `Banco: ${{DADOS.banco}} | Agência: ${{DADOS.agencia}} | Conta: ${{DADOS.contaCorrente}}`
    );

    console.log(
        `Número Título: ${{DADOS.numeroTitulo}}`
    );

    console.log(
        '=========================================='
    );

}})();
"""
    return PlainTextResponse(content=script, media_type="application/javascript")


@router.post("/{reconciliation_id}/fechar", response_model=ReconciliationOut)
def fechar_competencia(
    reconciliation_id: UUID, performed_by: str,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    reconciliation.status = ReconciliationLifecycleStatus.FECHADA
    reconciliation.closed_at = datetime.now(timezone.utc)
    reconciliation.closed_by = performed_by
    session.flush()

    session.add(AuditLog(
        action=AuditAction.FECHAMENTO, entity_type="Reconciliation", entity_id=reconciliation.id,
        performed_by=performed_by,
    ))
    session.commit()
    return reconciliation


@router.post("/{reconciliation_id}/reabrir", response_model=ReconciliationOut)
def reabrir_competencia(
    reconciliation_id: UUID, performed_by: str,
    session: Session = Depends(get_session), _=Depends(require_api_key),
):
    reconciliation = _get_reconciliation_or_404(session, reconciliation_id)
    reconciliation.status = ReconciliationLifecycleStatus.REABERTA
    session.flush()

    session.add(AuditLog(
        action=AuditAction.REABERTURA, entity_type="Reconciliation", entity_id=reconciliation.id,
        performed_by=performed_by,
    ))
    session.commit()
    return reconciliation
